"""Editable DOCX exports that closely reproduce the existing Harvard/Executive designs.

The PDF is produced by Chromium from HTML/CSS. Word uses a different layout
engine, so exact pixel-for-pixel identity cannot be guaranteed while keeping
text, tables, and images editable. Both formats use the same structured CV.
"""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from resumeforge.cv_languages import HEADINGS
from resumeforge.model import flatten_skills, normalize_profile

NAVY=RGBColor(22,35,57)
BURGUNDY=RGBColor(166,28,53)
MUTED=RGBColor(83,100,122)
GRAY=RGBColor(64,70,77)
WHITE=RGBColor(248,250,252)
LIGHT=RGBColor(241,245,249)
BLUE=RGBColor(21,43,69)


def _txt(x):
    if x is None:return ''
    if isinstance(x,(tuple,list)):return ', '.join(map(_txt,x))
    return str(x)


def _shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd');tcPr.append(shd)
    shd.set(qn('w:fill'),color)


def _border(paragraph,side='top',color='D7DEE6',sz='6'):
    ppr=paragraph._p.get_or_add_pPr()
    bd=ppr.find(qn('w:pBdr'))
    if bd is None:bd=OxmlElement('w:pBdr');ppr.append(bd)
    edge=OxmlElement('w:'+side)
    for k,v in {'val':'single','sz':sz,'color':color,'space':'3'}.items():edge.set(qn('w:'+k),v)
    bd.append(edge)


def _cell_border(cell,edge,color='A61C35',sz='12'):
    tcpr=cell._tc.get_or_add_tcPr()
    bd=tcpr.find(qn('w:tcBorders'))
    if bd is None:bd=OxmlElement('w:tcBorders');tcpr.append(bd)
    el=OxmlElement('w:'+edge)
    for k,v in {'val':'single','sz':sz,'color':color}.items():el.set(qn('w:'+k),v)
    bd.append(el)


def _margins(cell, top=65, start=95, bottom=60, end=95):
    tcpr=cell._tc.get_or_add_tcPr()
    margins=tcpr.find(qn('w:tcMar'))
    if margins is None:margins=OxmlElement('w:tcMar');tcpr.append(margins)
    for name,value in {'top':top,'start':start,'bottom':bottom,'end':end}.items():
        e=OxmlElement('w:'+name);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)


def _keep(p,next_=False):
    p.paragraph_format.keep_together=True
    if next_:p.paragraph_format.keep_with_next=True


def _p(parent, text='', size=9, bold=False, color=NAVY, before=0, after=3, align=None, italic=False):
    p=parent.add_paragraph() if not hasattr(parent,'_tc') else parent.add_paragraph()
    p.paragraph_format.space_before=Pt(before)
    p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.line_spacing=1.25
    if align is not None:p.alignment=align
    r=p.add_run(_txt(text));r.bold=bold;r.italic=italic
    r.font.size=Pt(size);r.font.color.rgb=color
    return p


def _section(parent,title,color=BURGUNDY,before=8,after=5):
    p=_p(parent,title.upper(),size=10,bold=True,color=color,before=before,after=after)
    _border(p,'top', 'D7DEE6', '8')
    _keep(p,True)
    return p


def _bullet(parent,text,shade=False):
    p=_p(parent,_txt(text),size=8.8,after=2)
    p.style='List Bullet'
    p.paragraph_format.left_indent=Inches(.18)
    p.paragraph_format.first_line_indent=Inches(-.12)
    _keep(p)
    return p


def _no_break(row):
    trpr=row._tr.get_or_add_trPr();node=OxmlElement('w:cantSplit');trpr.append(node)


def _setup(doc, design,language):
    sec=doc.sections[0]
    if design=='harvard':
        sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.top_margin=Inches(.44);sec.bottom_margin=Inches(.5)
        sec.left_margin=Inches(.63);sec.right_margin=Inches(.63)
    else:
        sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
        sec.top_margin=Inches(.36);sec.bottom_margin=Inches(.38)
        sec.left_margin=Inches(.32);sec.right_margin=Inches(.34)
    sec.header_distance=Inches(.15)
    normal=doc.styles['Normal']
    name='Carlito' if language not in ('ja','zh') else ('Noto Sans CJK JP' if language=='ja' else 'Noto Sans CJK SC')
    normal.font.name=name;normal.font.size=Pt(9);normal.font.color.rgb=NAVY
    normal.paragraph_format.space_after=Pt(3)
    normal.paragraph_format.line_spacing=1.25
    if language in ('ja','zh'):
        rpr=normal.element.get_or_add_rPr();fonts=rpr.rFonts
        if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
        fonts.set(qn('w:eastAsia'),name)


def _harvard(doc,p,labels):
    title=_p(doc,_txt(p.get('name')).upper(),size=17.5,bold=True,color=NAVY,after=2,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    _keep(title,True)
    contact='  •  '.join(_txt(p.get(k)) for k in ('location','email','phone') if p.get(k))
    if contact:_p(doc,contact,size=8.9,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
    social='  •  '.join(_txt(p.get(k)) for k in ('linkedin','github') if p.get(k))
    if social:_p(doc,social,size=8.2,color=MUTED,align=WD_ALIGN_PARAGRAPH.CENTER,after=11)
    if p.get('summary'):
        _section(doc,labels['profile'],after=4)
        _p(doc,p['summary'],after=7)
    skills=flatten_skills(p)
    if skills:
        _section(doc,labels['skills'],after=5)
        cols=3 if len(skills)>4 else 2
        for n in range(0,len(skills),cols):
            t=doc.add_table(rows=1,cols=cols);t.autofit=False
            for j in range(cols):
                c=t.cell(0,j);c.width=Inches(2.25);_margins(c,top=35,start=65,bottom=35,end=65)
                if n+j<len(skills):
                    _shade(c,'F8FAFC')
                    for edge in ('top','bottom','start','end'):_cell_border(c,edge,'D7E2EB','4')
                    r=c.paragraphs[0].add_run(_txt(skills[n+j]));r.font.size=Pt(7.5);r.font.color.rgb=MUTED
                else:_shade(c,'FFFFFF')
            _p(doc,'',after=0)
    if p.get('experience'):
        _section(doc,labels['experience'],after=4)
        for job in p['experience']:
            if not isinstance(job,dict):continue
            tbl=doc.add_table(rows=1,cols=2);tbl.autofit=False
            tbl.columns[0].width=Inches(.09);tbl.columns[1].width=Inches(6.95)
            left,right=tbl.rows[0].cells
            _shade(left,'A61C35');_margins(left,0,0,0,0)
            _margins(right,top=20,start=125,bottom=30,end=35)
            _no_break(tbl.rows[0]);_keep(right.paragraphs[0],True)
            hp=right.paragraphs[0];hp.paragraph_format.space_after=Pt(2)
            r=hp.add_run(_txt(job.get('title')));r.bold=True;r.font.size=Pt(9.4);r.font.color.rgb=NAVY
            if job.get('period'):
                r=hp.add_run('    '+_txt(job['period']));r.italic=True;r.font.size=Pt(8.2);r.font.color.rgb=MUTED
            _p(right,job.get('organization'),size=8.5,bold=True,color=BURGUNDY,after=2)
            for bullet in job.get('bullets') or []:_bullet(right,bullet)
            if job.get('description'):_p(right,job['description'],size=8.6)
            _p(doc,'',after=1)
    if p.get('education'):
        _section(doc,labels['education'])
        for e in p['education']:
            if not isinstance(e,dict):continue
            x=_p(doc,e.get('qualification'),size=9,bold=True,after=0)
            _keep(x,True)
            _p(doc,' · '.join(_txt(e.get(k)) for k in ('institution','year') if e.get(k)),size=8.4,color=MUTED,after=3)
    if p.get('projects'):
        _section(doc,labels['projects'])
        proj=p['projects']
        for i in range(0,len(proj),2):
            tab=doc.add_table(rows=1,cols=2);tab.autofit=False
            for j in range(2):
                cell=tab.cell(0,j);_shade(cell,'F1F5F9');_cell_border(cell,'top','A61C35','8');_cell_border(cell,'bottom','A61C35','8')
                _margins(cell,55,100,70,100)
                if i+j<len(proj):
                    project=proj[i+j]
                    if not isinstance(project,dict):continue
                    r=cell.paragraphs[0].add_run(_txt(project.get('name')));r.font.color.rgb=RGBColor(52,86,116);r.font.size=Pt(9);r.bold=True
                    _p(cell,project.get('technologies'),size=8.2,color=MUTED,after=2)
                    _p(cell,project.get('description'),size=8.5,after=2)
            _p(doc,'',after=2)
    for key,label in [('training','training'),('languages','languages')]:
        if p.get(key):
            _section(doc,labels[label])
            _p(doc,'  |  '.join(map(_txt,p[key])),size=8.8)
    if p.get('extra_notes'):
        _section(doc,labels['notes']);_p(doc,p['extra_notes'],size=8.6)


def _sidebar(cell,p,index,photo,brain,labels):
    _shade(cell,'40464D');_margins(cell,145,125,160,125)
    if index==0:
        if photo:
            try:
                paragraph=cell.paragraphs[0]
                paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(BytesIO(photo),width=Inches(1.52),height=Inches(1.55))
            except Exception:pass
        else:
            first=_txt(p.get('name'))[:1].upper()
            z=_p(cell,first,size=28,bold=True,color=WHITE,after=12,align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        _p(cell,'',size=3,after=9)
    _p(cell,_txt(p.get('name')).upper(),size=12,bold=True,color=WHITE,after=3,
       align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(cell,p.get('headline') or 'TECHNOLOGY · OPERATIONS',size=8,color=WHITE,after=13,
       align=WD_ALIGN_PARAGRAPH.CENTER)
    if p.get('tagline'):
        _p(cell,p['tagline'],size=10,color=WHITE,after=14,align=WD_ALIGN_PARAGRAPH.CENTER)
    if index==0:
        for k in ('location','phone','email','linkedin','github'):
            if p.get(k):_p(cell,p[k],size=7.5,color=WHITE,after=11)
        if p.get('languages'):
            _p(cell,labels['languages'].upper(),size=8,bold=True,color=WHITE,before=9,after=9,
               align=WD_ALIGN_PARAGRAPH.CENTER)
            for lang in p['languages']:_p(cell,lang,size=8,color=WHITE,after=8)
    else:
        _p(cell,labels['overview'].upper(),size=8,bold=True,color=WHITE,before=22,after=9,
           align=WD_ALIGN_PARAGRAPH.CENTER)
        for group in p.get('skill_groups',[])[:4]:
            _p(cell,group.get('title',''),size=7.8,bold=True,color=WHITE,after=3)
            _p(cell,', '.join(map(_txt,group.get('items',[])[:3])),size=7.6,color=WHITE,after=11)
    if brain:
        try:
            _p(cell,'',after=10)
            r=cell.add_paragraph().add_run();r.add_picture(BytesIO(brain),width=Inches(1.5))
        except Exception:pass


def _executive_body(cell,blocks,p,labels):
    _margins(cell,200,100,80,90)
    for block in blocks:
        kind=block.get('type')
        if kind=='profile':
            _section(cell,labels['profile'],color=BLUE,before=0)
            _p(cell,p.get('summary',''),size=9.4,after=15)
        elif kind=='skills':
            _section(cell,labels['competencies'],color=BLUE)
            groups=p.get('skill_groups',[])[:4]
            if not groups:
                sk=flatten_skills(p)
                groups=[{'title':labels['core'],'items':sk[:8]},{'title':labels['technical'],'items':sk[8:16]},
                        {'title':labels['more'],'items':sk[16:24]}, {'title':labels['other'],'items':sk[24:]}]
            for i in range(0,len(groups),2):
                grid=cell.add_table(rows=1,cols=2)
                for j in range(2):
                    c=grid.cell(0,j);_shade(c,'F2F5F7');_margins(c,130,90,75,75)
                    if i+j<len(groups):
                        g=groups[i+j]
                        rr=c.paragraphs[0].add_run(_txt(g.get('title')).upper());rr.bold=True;rr.font.size=Pt(8.9);rr.font.color.rgb=BLUE
                        for item in g.get('items',[]):_p(c,'• '+_txt(item),size=8,color=BLUE,after=3)
        elif kind=='heading':
            key='continued' if 'Continued' in block.get('text','') else 'experience'
            _section(cell,labels[key],color=BLUE,before=9)
        elif kind=='role':
            j=block['job']
            h=_p(cell,j.get('title'),size=10,bold=True,color=BLUE,before=6,after=2)
            _keep(h,True)
            _p(cell,'  |  '.join(_txt(j.get(k)) for k in ('organization','period') if j.get(k)),
               size=8.5,color=MUTED,after=5)
            for bullet in j.get('bullets') or []:_bullet(cell,bullet)
        elif kind=='education':
            _section(cell,labels['education'],color=BLUE)
            for ed in p.get('education',[]):
                if isinstance(ed,dict):
                    _p(cell,ed.get('qualification'),size=9.8,bold=True,after=1)
                    _p(cell,' · '.join(_txt(ed.get(k)) for k in ('institution','year') if ed.get(k)),size=8.5,color=MUTED,after=8)
        elif kind=='projects':
            _section(cell,labels['projects'],color=BLUE)
            items=p.get('projects',[])
            for n in range(0,len(items),2):
                t=cell.add_table(rows=1,cols=2)
                for j in range(2):
                    c=t.cell(0,j);_shade(c,'F2F5F7');_margins(c,90,70,100,70)
                    if n+j<len(items):
                        pr=items[n+j]
                        _p(c,pr.get('name'),size=9.2,bold=True,color=BLUE,after=2)
                        _p(c,pr.get('technologies'),size=8,color=MUTED,after=2)
                        _p(c,pr.get('description'),size=8.5,after=3)
        elif kind in ('training','languages','notes'):
            kk={'training':'training','languages':'languages','notes':'notes'}[kind]
            value=p.get('extra_notes') if kind=='notes' else '; '.join(map(_txt,p.get(kk,[])))
            _section(cell,labels[kk],color=BLUE)
            _p(cell,value,size=8.5)


def _executive(doc,p,labels,photo=None,brain=None):
    try:
        from resumeforge.render import _pack_executive
        pages=_pack_executive(deepcopy(p))
    except Exception:
        pages=[{'index':0,'blocks':[{'type':'profile'},{'type':'heading','text':'Professional Experience'}] +
                [{'type':'role','job':j} for j in p.get('experience',[])] +
                ([{'type':'projects'}] if p.get('projects') else [])}]
    for i,page in enumerate(pages):
        if i:doc.add_page_break()
        table=doc.add_table(rows=1,cols=2)
        table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width=Inches(1.98);table.columns[1].width=Inches(5.38)
        left,right=table.rows[0].cells
        left.width=Inches(1.98);right.width=Inches(5.38)
        left.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        right.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        _no_break(table.rows[0])
        # Deliberately taller left panel than text so it visually extends down page.
        table.rows[0].height=Inches(10.15)
        table.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        _sidebar(left,p,i,photo if i==0 else None,brain,labels)
        _executive_body(right,page['blocks'],p,labels)


def build_docx(profile: dict, design: str = 'Harvard', language: str='en', *, photo=None, brain=None) -> bytes:
    """Return an editable DOCX with a much closer layout to the PDF selected template."""
    if not isinstance(profile,dict):raise TypeError('CV profile must be a dictionary.')
    if language not in HEADINGS:raise ValueError('Unsupported CV language')
    choice='harvard' if str(design).lower().startswith('harvard') else 'executive'
    cv=normalize_profile(deepcopy(profile))
    doc=Document();_setup(doc,choice,language)
    doc.core_properties.title=_txt(cv.get('name'))+' - ResumeForge CV'
    doc.core_properties.author='ResumeForge'
    if choice=='harvard':
        _harvard(doc,cv,HEADINGS[language])
    else:
        if brain is None:
            from pathlib import Path
            art=Path(__file__).resolve().parent.parent/'assets'/'ai_brain_art.png'
            if art.exists():brain=art.read_bytes()
        _executive(doc,cv,HEADINGS[language],photo,brain)
    out=BytesIO();doc.save(out)
    return out.getvalue()
